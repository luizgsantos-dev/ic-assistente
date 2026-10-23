from datetime import date

import docx

from app.ingestao.leitura import _titulo_util, dividir_markdown, ler_documento
from tests.pdfs import criar_pdf, criar_pdf_sem_texto


def test_markdown_divide_por_secao_com_caminho():
    texto = (
        "# Regimento\n\nIntrodução.\n\n## Capítulo I\n\nArt. 1 texto.\n\n"
        "### Seção 1\n\nArt. 2 texto.\n\n## Capítulo II\n\nArt. 3 texto."
    )
    blocos = dividir_markdown(texto)
    assert [(b.secao, b.texto) for b in blocos] == [
        ("Regimento", "Introdução."),
        ("Regimento > Capítulo I", "Art. 1 texto."),
        ("Regimento > Capítulo I > Seção 1", "Art. 2 texto."),
        ("Regimento > Capítulo II", "Art. 3 texto."),
    ]


def test_secao_continua_na_pagina_seguinte():
    pilha: list = []
    p1 = dividir_markdown("# Calendário\n\nInício das aulas.", pagina=1, pilha=pilha)
    p2 = dividir_markdown("Fim das aulas.", pagina=2, pilha=pilha)
    assert p1[0].secao == p2[0].secao == "Calendário"
    assert p2[0].pagina == 2


def test_titulo_com_negrito_e_vazio():
    blocos = dividir_markdown("## **Semana Acadêmica**\n\nTexto.\n\n#\n\nOutro.")
    assert blocos[0].secao == "Semana Acadêmica"
    # "#" sozinho não é título; vira texto comum.
    assert "Outro." in blocos[-1].texto


def test_pdf_com_paginas_titulo_e_data(tmp_path):
    caminho = criar_pdf(
        tmp_path / "cal.pdf",
        [
            [("Calendário Acadêmico 2026", 20), ("O semestre 2026/2 termina em 12/12/2026.", 11)],
            [("Semana Acadêmica", 20), ("Evento anual com palestras e minicursos.", 11)],
        ],
        titulo="Calendário 2026",
        data_criacao="D:20260115103000",
    )
    lido = ler_documento(caminho)
    assert not lido.escaneado
    assert lido.paginas == 2
    assert lido.titulo == "Calendário 2026"
    assert lido.data == date(2026, 1, 15)
    paginas = {b.pagina for b in lido.blocos}
    assert paginas == {1, 2}
    texto = " ".join(b.texto for b in lido.blocos)
    assert "termina em 12/12/2026" in texto
    assert any(b.secao and "Semana Acadêmica" in b.secao for b in lido.blocos)


def test_pdf_sem_texto_e_marcado_como_escaneado(tmp_path):
    lido = ler_documento(criar_pdf_sem_texto(tmp_path / "scan.pdf"))
    assert lido.escaneado
    assert lido.blocos == []


def test_docx_com_titulos_e_tabela(tmp_path):
    d = docx.Document()
    d.core_properties.title = "Guia de Aproveitamento"
    d.add_heading("Aproveitamento de Estudos", level=1)
    d.add_paragraph("O pedido é feito na secretaria.")
    tabela = d.add_table(rows=2, cols=2)
    tabela.cell(0, 0).text, tabela.cell(0, 1).text = "Documento", "Onde"
    tabela.cell(1, 0).text, tabela.cell(1, 1).text = "Histórico", "SIGAA"
    caminho = tmp_path / "guia.docx"
    d.save(caminho)

    lido = ler_documento(caminho)
    assert lido.titulo == "Guia de Aproveitamento"
    assert all(b.secao == "Aproveitamento de Estudos" for b in lido.blocos)
    texto = "\n".join(b.texto for b in lido.blocos)
    assert "O pedido é feito na secretaria." in texto
    assert "| Histórico | SIGAA |" in texto


def test_txt_em_windows_1252(tmp_path):
    caminho = tmp_path / "aviso.txt"
    caminho.write_bytes("Inscrições até sexta.".encode("cp1252"))
    assert ler_documento(caminho).blocos[0].texto == "Inscrições até sexta."


def test_titulo_dos_metadados():
    assert _titulo_util("Microsoft Word - Resolucao 12-2025.docx") == "Resolucao 12-2025"
    assert _titulo_util("untitled") is None
    assert _titulo_util(None) is None
